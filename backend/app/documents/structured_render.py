"""Real structured-content rendering for Word/PowerPoint/Excel documents.

Unlike app/documents/render.py (which turns Claude-drafted markdown into
a file), these renderers take fully structured content supplied by the
caller and need no AI provider at all — python-docx, python-pptx, and
openpyxl each produce a genuine, openable binary file directly. Never a
placeholder/renamed-.txt file: every output here is real docx/pptx/xlsx.
"""
import io

from app.schemas.generation import ExcelDocumentRequest, PptDocumentRequest, WordDocumentRequest


def render_structured_docx(payload: WordDocumentRequest) -> bytes:
    from docx import Document as DocxDocument
    from docx.shared import Pt

    doc = DocxDocument()
    if payload.author:
        doc.core_properties.author = payload.author
    doc.add_heading(payload.title, level=0)

    for block in payload.blocks:
        if block.type == "heading":
            doc.add_heading(block.text or "", level=min(block.level, 4))
        elif block.type == "bullet_list":
            for item in block.items or []:
                doc.add_paragraph(item, style="List Bullet")
        elif block.type == "numbered_list":
            for item in block.items or []:
                doc.add_paragraph(item, style="List Number")
        elif block.type == "table":
            rows = block.rows or []
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r, row_values in enumerate(rows):
                for c, cell_value in enumerate(row_values):
                    cell = table.cell(r, c)
                    cell.text = str(cell_value)
                    if block.header_row and r == 0:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
        else:  # paragraph
            p = doc.add_paragraph()
            run = p.add_run(block.text or "")
            run.bold = block.bold
            run.italic = block.italic
            run.font.size = Pt(11)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


_LAYOUT_INDEX = {
    "title": 0,
    "title_content": 1,
    "section_header": 2,
    "blank": 6,
}


def render_pptx(payload: PptDocumentRequest) -> bytes:
    from pptx import Presentation

    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = payload.title
    if payload.subtitle and len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = payload.subtitle

    for slide_spec in payload.slides:
        layout_index = _LAYOUT_INDEX.get(slide_spec.layout, 1)
        layout = prs.slide_layouts[layout_index]
        slide = prs.slides.add_slide(layout)

        if slide.shapes.title is not None:
            slide.shapes.title.text = slide_spec.title

        if slide_spec.layout == "title" and slide_spec.subtitle and len(slide.placeholders) > 1:
            slide.placeholders[1].text = slide_spec.subtitle
        elif slide_spec.bullets and len(slide.placeholders) > 1:
            body = slide.placeholders[1].text_frame
            body.clear()
            for i, bullet in enumerate(slide_spec.bullets):
                paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
                paragraph.text = bullet
                paragraph.level = 0

        if slide_spec.notes:
            slide.notes_slide.notes_text_frame.text = slide_spec.notes

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def render_xlsx(payload: ExcelDocumentRequest) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_spec in payload.sheets:
        ws = wb.create_sheet(title=sheet_spec.name)

        ws.append(sheet_spec.headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for row in sheet_spec.rows:
            ws.append(row)

        for col_index, header in enumerate(sheet_spec.headers, start=1):
            max_len = max([len(str(header))] + [len(str(row[col_index - 1])) for row in sheet_spec.rows if col_index - 1 < len(row)])
            ws.column_dimensions[get_column_letter(col_index)].width = min(max_len + 4, 60)

        if sheet_spec.freeze_header:
            ws.freeze_panes = "A2"

        header_index = {name: idx for idx, name in enumerate(sheet_spec.headers, start=1)}
        n_rows = len(sheet_spec.rows)
        chart_anchor_col = len(sheet_spec.headers) + 2

        for chart_spec in sheet_spec.charts:
            if chart_spec.category_column not in header_index or n_rows == 0:
                continue
            chart_cls = {"bar": BarChart, "line": LineChart, "pie": PieChart}[chart_spec.type]
            chart = chart_cls()
            chart.title = chart_spec.title or chart_spec.type.capitalize()

            cat_col = header_index[chart_spec.category_column]
            categories = Reference(ws, min_col=cat_col, min_row=2, max_row=n_rows + 1)

            for value_column in chart_spec.value_columns:
                if value_column not in header_index:
                    continue
                val_col = header_index[value_column]
                data = Reference(ws, min_col=val_col, min_row=1, max_row=n_rows + 1)
                chart.add_data(data, titles_from_data=True)

            chart.set_categories(categories)
            ws.add_chart(chart, f"{get_column_letter(chart_anchor_col)}2")
            chart_anchor_col += 10

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
